"""Command-line interface for Ready for AI."""

import os
import sys
import click
from typing import Optional

from .detectors.pii_detector import PIIDetector
from .detectors.patterns import PIIType, PIIMatch
from .storage.mapping_store import MappingStore
from .storage.learning_store import LearningStore
from .processors.docx_processor import DocxProcessor, DocxRestorer
from .processors.pdf_processor import PdfProcessor, PdfRestorer


def get_user_decision(match: PIIMatch) -> Optional[bool]:
    """
    Prompt user for decision on uncertain PII detection.

    Returns:
        True if PII, False if not PII, None to skip
    """
    click.echo()
    click.echo(click.style("Uncertain detection found:", fg="yellow", bold=True))
    click.echo(f"  Type: {match.pii_type.value}")
    click.echo(f"  Value: {click.style(match.text, fg='cyan', bold=True)}")
    click.echo(f"  Confidence: {match.confidence:.0%}")
    click.echo(f"  Context: ...{match.context}...")
    click.echo()

    while True:
        choice = click.prompt(
            "Is this confidential/PII? [y]es, [n]o, [s]kip",
            type=str,
            default="s"
        ).lower()

        if choice in ('y', 'yes'):
            # Ask for correct type if user confirms
            click.echo("Confirm PII type:")
            for i, pii_type in enumerate(PIIType, 1):
                click.echo(f"  {i}. {pii_type.value}")
            click.echo(f"  0. Keep as {match.pii_type.value}")

            type_choice = click.prompt("Type number", type=int, default=0)
            if type_choice > 0 and type_choice <= len(PIIType):
                match.pii_type = list(PIIType)[type_choice - 1]

            return True
        elif choice in ('n', 'no'):
            return False
        elif choice in ('s', 'skip'):
            return None
        else:
            click.echo("Please enter 'y', 'n', or 's'")


@click.group()
@click.version_option(version="0.1.0")
def cli():
    """Ready for AI - Prepare documents for AI by redacting PII."""
    pass


@cli.command()
@click.argument('input_file', type=click.Path(exists=True))
@click.option('-o', '--output', type=click.Path(), help='Output file path')
@click.option('-m', '--mapping-file', type=click.Path(),
              help='Path to save/load mapping file')
@click.option('-p', '--password', prompt=True, hide_input=True,
              confirmation_prompt=True,
              help='Password for encrypting mappings')
@click.option('--no-interactive', is_flag=True,
              help='Disable interactive mode (skip uncertain detections)')
@click.option('--redact-all', is_flag=True,
              help='Redact all detections including uncertain ones (no prompts)')
@click.option('--no-nlp', is_flag=True,
              help='Disable NLP-based detection (faster but less accurate)')
def redact(
    input_file: str,
    output: Optional[str],
    mapping_file: Optional[str],
    password: str,
    no_interactive: bool,
    redact_all: bool,
    no_nlp: bool,
):
    """
    Redact PII from a document.

    Supported formats: DOCX, PDF
    """
    # Convert to absolute path
    input_file = os.path.abspath(input_file)

    # Determine file type
    ext = os.path.splitext(input_file)[1].lower()
    if ext not in ('.docx', '.pdf'):
        click.echo(click.style(f"Unsupported file type: {ext}", fg="red"))
        click.echo("Supported formats: .docx, .pdf")
        sys.exit(1)

    # Set default mapping file
    if mapping_file is None:
        base = os.path.splitext(input_file)[0]
        mapping_file = f"{base}_mappings.json"

    click.echo(click.style("Ready for AI - PII Redaction Tool", fg="green", bold=True))
    click.echo()

    # Initialize components
    click.echo("Initializing...")
    learning_store = LearningStore()
    learned_data = learning_store.get_learned_data()

    detector = PIIDetector(
        use_nlp=not no_nlp,
        learned_patterns={k: set(v) for k, v in learned_data['pii'].items()},
        learned_safe=set(learned_data['safe']),
    )

    mapping_store = MappingStore(password=password)

    # Set up callback for learning
    def learning_callback(match: PIIMatch) -> Optional[bool]:
        decision = get_user_decision(match)
        if decision is True:
            learning_store.learn_pii(match.text, match.pii_type.value)
        elif decision is False:
            learning_store.learn_safe(match.text)
        return decision

    # redact_all: always return True for uncertain items
    def redact_all_callback(match: PIIMatch) -> Optional[bool]:
        return True

    if redact_all:
        callback = redact_all_callback
        interactive = True  # Need interactive=True so callback is called
    elif no_interactive:
        callback = None
        interactive = False
    else:
        callback = learning_callback
        interactive = True

    # Process based on file type
    click.echo(f"Processing: {input_file}")

    if ext == '.docx':
        processor = DocxProcessor(
            detector=detector,
            mapping_store=mapping_store,
            interactive=interactive,
            user_callback=callback,
        )
        result = processor.process(input_file, output)
    else:  # PDF
        processor = PdfProcessor(
            detector=detector,
            mapping_store=mapping_store,
            interactive=interactive,
            user_callback=callback,
        )
        result = processor.process(input_file, output)

    # Save mapping file
    mapping_store.save_to_file(mapping_file)

    # Show results
    click.echo()
    click.echo(click.style("Redaction Complete!", fg="green", bold=True))
    click.echo(f"  Input:  {result.input_path}")
    click.echo(f"  Output: {result.output_path}")
    click.echo(f"  Mapping file: {mapping_file}")
    click.echo()
    click.echo(f"  Total redactions: {result.total_redactions}")
    click.echo(f"  Uncertain items: {result.uncertain_count}")

    if result.redactions_by_type:
        click.echo("  By type:")
        for pii_type, count in sorted(result.redactions_by_type.items()):
            click.echo(f"    - {pii_type}: {count}")

    click.echo()
    click.echo(click.style(
        "Keep the mapping file safe! You'll need it and the password to restore.",
        fg="yellow"
    ))


@cli.command()
@click.argument('input_file', type=click.Path(exists=True))
@click.argument('mapping_file', type=click.Path(exists=True))
@click.option('-o', '--output', type=click.Path(), help='Output file path')
@click.option('-p', '--password', prompt=True, hide_input=True,
              help='Password used when creating mappings')
@click.option('--delete-mapping', is_flag=True,
              help='Delete mapping file after successful restoration')
def restore(
    input_file: str,
    mapping_file: str,
    output: Optional[str],
    password: str,
    delete_mapping: bool,
):
    """
    Restore original values in a redacted document.

    Requires the mapping file created during redaction and the password.
    """
    # Convert to absolute paths
    input_file = os.path.abspath(input_file)
    mapping_file = os.path.abspath(mapping_file)

    ext = os.path.splitext(input_file)[1].lower()
    if ext not in ('.docx', '.pdf'):
        click.echo(click.style(f"Unsupported file type: {ext}", fg="red"))
        sys.exit(1)

    click.echo(click.style("Ready for AI - Restoration Tool", fg="green", bold=True))
    click.echo()

    # Load mapping store
    click.echo("Loading mappings...")
    try:
        mapping_store = MappingStore.load_from_file(mapping_file, password)
    except Exception as e:
        click.echo(click.style(f"Error loading mappings: {e}", fg="red"))
        click.echo("Check that the password is correct and the file is valid.")
        sys.exit(1)

    # Restore
    click.echo(f"Restoring: {input_file}")

    try:
        if ext == '.docx':
            restorer = DocxRestorer(mapping_store)
            output_path, count = restorer.restore(input_file, output)
        else:  # PDF
            restorer = PdfRestorer(mapping_store)
            output_path, count = restorer.restore(input_file, output)
    except Exception as e:
        click.echo(click.style(f"Error during restoration: {e}", fg="red"))
        sys.exit(1)

    # Optionally delete mapping file
    if delete_mapping:
        os.remove(mapping_file)
        click.echo(f"Deleted mapping file: {mapping_file}")

    click.echo()
    click.echo(click.style("Restoration Complete!", fg="green", bold=True))
    click.echo(f"  Input:  {input_file}")
    click.echo(f"  Output: {output_path}")
    click.echo(f"  Restorations: {count}")


@cli.command()
@click.argument('input_file', type=click.Path(exists=True))
@click.option('--no-nlp', is_flag=True, help='Disable NLP-based detection')
def scan(input_file: str, no_nlp: bool):
    """
    Scan a document for PII without redacting.

    Shows what would be detected and allows reviewing/learning.
    """
    # Convert to absolute path
    input_file = os.path.abspath(input_file)

    ext = os.path.splitext(input_file)[1].lower()
    if ext not in ('.docx', '.pdf'):
        click.echo(click.style(f"Unsupported file type: {ext}", fg="red"))
        sys.exit(1)

    click.echo(click.style("Ready for AI - PII Scanner", fg="green", bold=True))
    click.echo()

    # Initialize
    learning_store = LearningStore()
    learned_data = learning_store.get_learned_data()

    detector = PIIDetector(
        use_nlp=not no_nlp,
        learned_patterns={k: set(v) for k, v in learned_data['pii'].items()},
        learned_safe=set(learned_data['safe']),
    )

    # Extract text
    click.echo(f"Scanning: {input_file}")

    if ext == '.docx':
        from docx import Document
        doc = Document(input_file)
        text = "\n".join(p.text for p in doc.paragraphs)
    else:  # PDF
        import pdfplumber
        with pdfplumber.open(input_file) as pdf:
            text = "\n".join(
                page.extract_text() or "" for page in pdf.pages
            )

    # Detect
    result = detector.detect(text)

    # Show results
    click.echo()
    click.echo(click.style("High-confidence detections:", fg="green", bold=True))
    if result.matches:
        for match in result.matches:
            click.echo(f"  [{match.pii_type.value}] {match.text} ({match.confidence:.0%})")
    else:
        click.echo("  None found")

    click.echo()
    click.echo(click.style("Uncertain detections:", fg="yellow", bold=True))
    if result.uncertain:
        for match in result.uncertain:
            click.echo(f"  [{match.pii_type.value}] {match.text} ({match.confidence:.0%})")
    else:
        click.echo("  None found")

    # Learning mode
    click.echo()
    if result.uncertain and click.confirm("Would you like to review uncertain detections?"):
        for match in result.uncertain:
            decision = get_user_decision(match)
            if decision is True:
                learning_store.learn_pii(match.text, match.pii_type.value)
                click.echo(click.style(f"  Learned as PII: {match.text}", fg="green"))
            elif decision is False:
                learning_store.learn_safe(match.text)
                click.echo(click.style(f"  Learned as safe: {match.text}", fg="blue"))

    # Summary
    click.echo()
    click.echo(f"Total detections: {len(result.matches) + len(result.uncertain)}")
    click.echo(f"  High confidence: {len(result.matches)}")
    click.echo(f"  Uncertain: {len(result.uncertain)}")


@cli.command()
def learn_stats():
    """Show statistics about learned patterns."""
    learning_store = LearningStore()
    stats = learning_store.stats()

    click.echo(click.style("Learned Patterns Statistics", fg="green", bold=True))
    click.echo()
    click.echo(f"Total PII values learned: {stats['total_pii_values']}")
    click.echo(f"Total safe values learned: {stats['total_safe_values']}")
    click.echo(f"Custom patterns: {stats['custom_patterns']}")

    if stats['pii_by_type']:
        click.echo()
        click.echo("PII values by type:")
        for pii_type, count in sorted(stats['pii_by_type'].items()):
            click.echo(f"  - {pii_type}: {count}")


@cli.command()
@click.option('--export', type=click.Path(),
              help='Export learned patterns to file')
@click.option('--import-file', 'import_file', type=click.Path(exists=True),
              help='Import learned patterns from file')
@click.option('--clear', is_flag=True, help='Clear all learned patterns')
def learn_manage(export: Optional[str], import_file: Optional[str], clear: bool):
    """Manage learned patterns."""
    learning_store = LearningStore()

    if clear:
        if click.confirm("Are you sure you want to clear all learned patterns?"):
            learning_store.clear()
            click.echo("Learned patterns cleared.")
        return

    if export:
        learning_store.export_to_file(export)
        click.echo(f"Exported learned patterns to: {export}")
        return

    if import_file:
        merge = click.confirm("Merge with existing patterns? (No = replace)")
        learning_store.import_from_file(import_file, merge=merge)
        click.echo(f"Imported patterns from: {import_file}")
        return

    click.echo("Use --export, --import-file, or --clear")


@cli.command()
@click.argument('name')
@click.argument('pattern')
@click.option('-t', '--type', 'pii_type', default='custom',
              help='PII type for this pattern')
def add_pattern(name: str, pattern: str, pii_type: str):
    """Add a custom regex pattern for detection."""
    learning_store = LearningStore()
    learning_store.add_custom_pattern(name, pattern, pii_type)
    click.echo(f"Added custom pattern '{name}': {pattern}")


@cli.command()
@click.argument('mapping_file', type=click.Path(exists=True))
@click.option('-p', '--password', prompt=True, hide_input=True,
              help='Password used when creating mappings')
@click.option('-i', '--input-file', type=click.Path(exists=True),
              help='Read text from file instead of stdin')
@click.option('-o', '--output-file', type=click.Path(),
              help='Write restored text to file instead of stdout')
def restore_text(
    mapping_file: str,
    password: str,
    input_file: Optional[str],
    output_file: Optional[str],
):
    """
    Restore placeholders in any text (e.g., AI response).

    Reads text from stdin (or --input-file), replaces all placeholders
    with original values using the mapping file, outputs to stdout (or --output-file).

    Example workflow:
      1. ready-for-ai redact document.docx -p  # Creates redacted doc + mappings
      2. Send redacted doc to AI, get response with "John Doe", "Example Corp", etc.
      3. echo "AI response text" | ready-for-ai restore-text mappings.json -p
         OR: ready-for-ai restore-text mappings.json -p -i ai_response.txt -o restored.txt
    """
    # Convert to absolute paths
    mapping_file = os.path.abspath(mapping_file)
    if input_file:
        input_file = os.path.abspath(input_file)

    # Load mapping store
    try:
        mapping_store = MappingStore.load_from_file(mapping_file, password)
    except Exception as e:
        click.echo(click.style(f"Error loading mappings: {e}", fg="red"), err=True)
        sys.exit(1)

    # Get restorations
    restorations = mapping_store.get_all_restorations()
    if not restorations:
        click.echo(click.style("No mappings available for restoration", fg="red"), err=True)
        sys.exit(1)

    # Read input text
    if input_file:
        with open(input_file, 'r') as f:
            text = f.read()
    else:
        # Read from stdin
        if sys.stdin.isatty():
            click.echo("Enter text to restore (Ctrl+D to finish):", err=True)
        text = sys.stdin.read()

    # Apply restorations
    restored_text = text
    restoration_count = 0
    for placeholder, original in restorations.items():
        if placeholder in restored_text:
            count = restored_text.count(placeholder)
            restored_text = restored_text.replace(placeholder, original)
            restoration_count += count

    # Output
    if output_file:
        with open(output_file, 'w') as f:
            f.write(restored_text)
        click.echo(f"Restored {restoration_count} placeholders to: {output_file}", err=True)
    else:
        click.echo(restored_text)
        if sys.stderr.isatty():
            click.echo(f"\n--- Restored {restoration_count} placeholders ---", err=True)


@cli.command()
@click.option('-p', '--port', default=5000, help='Port to run the server on')
@click.option('--host', default='127.0.0.1', help='Host to bind to')
@click.option('--no-browser', is_flag=True, help='Do not open browser automatically')
def web(port: int, host: str, no_browser: bool):
    """
    Launch the web interface.

    Opens a browser with a user-friendly UI for redacting documents
    and restoring AI responses. No password required - session-based.
    """
    from .web import run_server

    click.echo(click.style("Ready for AI - Web Interface", fg="green", bold=True))
    run_server(host=host, port=port, open_browser=not no_browser)


def main():
    """Entry point."""
    cli()


if __name__ == '__main__':
    main()
